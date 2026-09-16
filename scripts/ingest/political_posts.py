# -*- coding: utf-8 -*-
"""Extract the three machine-readable political-appointee disclosures.

These are pay ENTITLEMENTS attached to posts - what a post pays - and they are
the mirror image of the premium data, which records payments that were made.
Nothing here is money anyone received: the Finance sheet has a no-pay-leave
column and termination dates, and four of the Foreign Affairs rows are blank
in every column, which is what a vacant post looks like.

Each CSV mirrors its own document's columns. They are not unified here; that
happens in build_graph.py. A single generic posts CSV would resemble none of
the three documents, and then nobody could check a figure by holding the page
beside the row, which is the entire reason these files exist.

Three documents, three shapes:

  Foreign Affairs   268 rows, whole ministry, as at 31 Dec 2025. 125 rows are
                    job type "Political"; the other 143 are the career Foreign
                    Service and are NOT political appointees.
  Health, Family    81 rows, as at 31 Mar 2026. No names - an RCN staff id.
   and Welfare      The rows are labelled with the pre-merger office name.
  Finance           34 rows, the only document that names anybody.

Run:  python scripts/ingest/political_posts.py [--accept]
Out:  data/political-posts-foreign-affairs.csv
      data/political-posts-health.csv
      data/political-posts-finance.csv
"""
import os
import re
import sys

import openpyxl
import pdfplumber

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import csvio  # noqa: E402

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, '..', '..'))
DOCS = os.path.join(ROOT, 'public', 'sources', 'political-appointees')

FOREIGN_PDF = os.path.join(DOCS, 'foreign-affairs-salary-structure-2025-12-31.pdf')
HEALTH_PDF = os.path.join(
    DOCS, 'health-family-welfare-political-appointees-2026-03-31.pdf')
FINANCE_XLSX = os.path.join(DOCS, 'finance-political-staff-list.xlsx')

FOREIGN_SOURCE = 'rti-foreign-affairs-salary-structure-2025'
HEALTH_SOURCE = 'rti-health-family-welfare-political-2026'
FINANCE_SOURCE = 'rti-finance-political-staff'

ROW_NO = re.compile(r'^\d+$')

# The designations, normalised onto the ladder. The source spellings are kept
# verbatim in the CSV - including the Health document's two typos - because a
# typo in a government pay disclosure is a fact about the document, and hiding
# it behind a lookup makes the CSV disagree with the page it came from.
RANKS = {
    'minister of foreign affairs': 'minister',
    'minister of finance and plaining': 'minister',
    'state minister': 'state-minister',
    'minister of state for foreign affairs': 'state-minister',
    'minister of state for finance and planning': 'state-minister',
    'minister of state for finance and planning (planning)': 'state-minister',
    'deputy minister': 'deputy-minister',
    'deputy minister (planning)': 'deputy-minister',
    'senior political director': 'senior-political-director',
    'senior political director (planning)': 'senior-political-director',
    'senior political dirctor': 'senior-political-director',     # sic
    'senier political director': 'senior-political-director',    # sic
    'political director': 'political-director',
}


def rank_for(designation):
    return RANKS.get((designation or '').strip().lower(), '')


# ---------------------------------------------------------------------------
# Ministry of Foreign Affairs
# ---------------------------------------------------------------------------

FOREIGN_COLUMNS = [
    'row_id', 'source_page', 'source_row', 'designation', 'rank', 'job_type',
    'basic_salary', 'foreign_service_allowance', 'living_allowance',
    'executive_allowance', 'service_allowance', 'supporting_co_allowance',
    'technical_co_allowance', 'phone_allowance', 'minimum_wage_allowance',
    'petrol_allowance', 'dress_allowance_yearly', 'stated_total',
]


def foreign_affairs():
    """The whole ministry's pay table, 268 rows across 7 pages.

    A `-` in the document is an explicit nil and becomes 0. A blank cell stays
    blank: 7 rows print no total at all, and 4 of those print nothing in any
    column. Those two are different facts and are never merged.
    """
    rows = []
    with pdfplumber.open(FOREIGN_PDF) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            for table in page.extract_tables():
                for cells in table:
                    if not cells or not ROW_NO.match((cells[0] or '').strip()):
                        continue
                    cells = [(c or '').strip() for c in cells]
                    designation = cells[1]
                    row = {
                        'row_id': f'fa--{int(cells[0]):03d}',
                        'source_page': pno,
                        'source_row': int(cells[0]),
                        'designation': designation,
                        'rank': rank_for(designation),
                        'job_type': cells[2],
                        # The document prints this to two places while
                        # printing its components rounded, so it is the
                        # only exact figure on the row. Never rounded.
                        'stated_total': csvio.amount(cells[14], dp=2),
                    }
                    for column, cell in zip(FOREIGN_COLUMNS[6:17], cells[3:14]):
                        row[column] = nil_or_blank(cell)
                    rows.append(row)
    return rows


def nil_or_blank(cell):
    """`-` is a printed nil and becomes 0; an empty cell stays empty."""
    cell = (cell or '').strip()
    if cell == '-':
        return '0'
    return csvio.amount(cell)


# ---------------------------------------------------------------------------
# Ministry of Health, Family and Welfare
# ---------------------------------------------------------------------------

HEALTH_COLUMNS = [
    'row_id', 'source_page', 'source_row', 'rcn', 'office', 'designation',
    'rank', 'employed_date', 'appraisal_marks_2025', 'department',
    'basic_salary', 'housing', 'transport', 'phone', 'other',
]

MONTHS = {m: i for i, m in enumerate(
    ['Jan', 'Feb', 'Mar', 'Apr', 'May', 'Jun',
     'Jul', 'Aug', 'Sep', 'Oct', 'Nov', 'Dec'], 1)}


def iso_date(text):
    """'29-Nov-23' -> '2023-11-29'. The document prints two-digit years."""
    match = re.fullmatch(r'(\d{1,2})-([A-Za-z]{3})-(\d{2})', (text or '').strip())
    if not match:
        return ''
    day, month, year = match.groups()
    return f'20{year}-{MONTHS[month.title()]:02d}-{int(day):02d}'


def health():
    """81 rows, no names.

    `office` is printed verbatim. The rows say "Ministry of Social and Family
    Development", the pre-merger name, while the responding body is Health,
    Family and Welfare. That divergence is itself a fact and is never
    corrected here.

    The Roles and Responsibilities column is deliberately not extracted: it is
    a tall wrapped cell whose text interleaves across band boundaries, and
    nothing renders it.
    """
    rows = []
    with pdfplumber.open(HEALTH_PDF) as pdf:
        for pno, page in enumerate(pdf.pages, 1):
            for table in page.extract_tables():
                for cells in table:
                    if not cells or not ROW_NO.match((cells[0] or '').strip()):
                        continue
                    cells = [(c or '').strip() for c in cells]
                    designation = cells[3]
                    rows.append({
                        'row_id': f'hfw--{int(cells[0]):03d}',
                        'source_page': pno,
                        'source_row': int(cells[0]),
                        'rcn': cells[1],
                        'office': cells[2],
                        'designation': designation,
                        'rank': rank_for(designation),
                        'employed_date': iso_date(cells[4]),
                        'appraisal_marks_2025': cells[5],
                        'department': cells[7],
                        'basic_salary': csvio.amount(cells[8]),
                        'housing': csvio.amount(cells[9]),
                        'transport': csvio.amount(cells[10]),
                        'phone': csvio.amount(cells[11]),
                        'other': csvio.amount(cells[12]),
                    })
    return rows


# ---------------------------------------------------------------------------
# Ministry of Finance and Planning
# ---------------------------------------------------------------------------

FINANCE_COLUMNS = [
    'row_id', 'source_row', 'name', 'office', 'designation', 'rank',
    'occupied_date', 'termination_date', 'rejoined_date', 'no_pay_leave',
    'basic_salary', 'basic_after_10pct_deduction', 'living_allowance',
    'phone_allowance', 'car_allowance',
]


def finance():
    """34 named appointees. The only document of the six that names anybody.

    `basic_after_10pct_deduction` is the sheet's own column: a 10 per cent cut
    applied to serving appointees until 31 December 2026. It is populated only
    for rows with no termination date, which is the sheet's way of saying the
    cut applies to whoever is still in post.
    """
    book = openpyxl.load_workbook(FINANCE_XLSX, data_only=True)
    sheet = book['Sheet1']
    rows = []
    for cells in sheet.iter_rows(min_row=4, max_row=sheet.max_row, values_only=True):
        if cells[0] is None or not str(cells[0]).strip().isdigit():
            continue
        designation = str(cells[3] or '').strip()
        rows.append({
            'row_id': f'fin--{int(cells[0]):03d}',
            'source_row': int(cells[0]),
            'name': str(cells[1] or '').strip(),
            'office': str(cells[2] or '').strip(),
            'designation': designation,
            'rank': rank_for(designation),
            'occupied_date': date_cell(cells[4]),
            'termination_date': date_cell(cells[5]),
            'rejoined_date': date_cell(cells[6]),
            'no_pay_leave': date_cell(cells[7]),
            'basic_salary': number_cell(cells[8]),
            'basic_after_10pct_deduction': number_cell(cells[9]),
            'living_allowance': number_cell(cells[10]),
            'phone_allowance': number_cell(cells[11]),
            'car_allowance': number_cell(cells[12]),
        })
    return rows


def date_cell(value):
    return value.strftime('%Y-%m-%d') if hasattr(value, 'strftime') else ''


def number_cell(value):
    return '' if value is None else csvio.amount(str(value))


def main():
    accept = '--accept' in sys.argv

    foreign = foreign_affairs()
    for row in foreign:
        row['source_id'] = FOREIGN_SOURCE
    csvio.sync('political-posts-foreign-affairs.csv',
               FOREIGN_COLUMNS + ['source_id'], foreign, accept)
    political = sum(1 for r in foreign if r['job_type'] == 'Political')
    print(f'  rows              {len(foreign)}')
    print(f'  political posts   {political}')
    print(f'  foreign service   {len(foreign) - political}')

    rows = health()
    for row in rows:
        row['source_id'] = HEALTH_SOURCE
    csvio.sync('political-posts-health.csv', HEALTH_COLUMNS + ['source_id'],
               rows, accept)
    print(f'  rows              {len(rows)}')

    rows = finance()
    for row in rows:
        row['source_id'] = FINANCE_SOURCE
    csvio.sync('political-posts-finance.csv', FINANCE_COLUMNS + ['source_id'],
               rows, accept)
    serving = sum(1 for r in rows if not r['termination_date'])
    print(f'  rows              {len(rows)}')
    print(f'  still in post     {serving}')


if __name__ == '__main__':
    main()
